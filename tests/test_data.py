"""Ingestion contract: edge-case fixtures plus, when present, the real raw workbook."""

from __future__ import annotations

import hashlib
import json
import os
import stat
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from restaurant_recs.data import (
    CATALOG_HEADERS, REVIEW_HEADERS, SECRET_FILENAME, DataValidationError,
    load_prepared, prepare,
)


def source_workbook() -> Path:
    configured = os.environ.get("RESTAURANT_WORKBOOK")
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parents[1] / "data" / "raw" / "RestaurantReviews.xlsx"


def make_workbook(path: Path, invalid_rating: bool = False) -> None:
    workbook = Workbook()
    catalog = workbook.active
    catalog.title = "Restaurants"
    catalog.append(CATALOG_HEADERS)
    catalog.append(("Claire's Korner", "Comfort food", 42.0, -87.7, 15, "Yes", "Local cafe"))
    reviews = workbook.create_sheet("Reviews")
    reviews.append(REVIEW_HEADERS)
    first = ("PRIVATE_REVIEWER_SENTINEL", "Clare's Korner", 5,
             "PRIVATE_REVIEW_TEXT_SENTINEL", datetime(2022, 1, 5), 1999,
             "PRIVATE_DEMOGRAPHIC_SENTINEL", True, False, 180, 70, 15, "Walk", True)
    reviews.append(first)
    reviews.append(first)  # Only this exact full-row duplicate should disappear.
    changed_review = list(first)
    changed_review[3] = "DIFFERENT_PRIVATE_REVIEW_SENTINEL"
    reviews.append(changed_review)  # Same projected event but distinct source row.
    unknown = list(first)
    unknown[1], unknown[2], unknown[4] = "Unlisted Place", 3, "4-0-2022"
    reviews.append(unknown)
    missing_date = list(first)
    missing_date[0], missing_date[4] = "SECOND_PRIVATE_REVIEWER", None
    reviews.append(missing_date)
    if invalid_rating:
        bad = list(first)
        bad[2] = 6
        reviews.append(bad)
    derived = workbook.create_sheet("north")
    derived.append(REVIEW_HEADERS)
    derived.append(first)
    workbook.save(path)
    workbook.close()


class SmallWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.workbook = self.root / "fixture.xlsx"
        make_workbook(self.workbook)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_exact_duplicate_policy_and_ignored_derived_sheet(self) -> None:
        prepared = prepare(self.workbook, self.root / "prepared")
        self.assertEqual(prepared["quality"]["raw_review_rows"], 5)
        self.assertEqual(prepared["quality"]["exact_duplicate_rows_removed"], 1)
        self.assertEqual(len(prepared["interactions"]), 4)
        self.assertEqual(prepared["quality"]["unique_user_item_pairs"], 3)
        self.assertEqual(prepared["quality"]["repeated_user_item_extra_rows"], 1)
        self.assertEqual(prepared["quality"]["matrix_observation_density"], 3 / (2 * 2))
        self.assertEqual(len({event["event_id"] for event in prepared["interactions"]}), 4)
        self.assertEqual(prepared["provenance"]["worksheets_ignored"], ["north"])

    def test_alias_unknown_item_and_bad_dates_remain_explicit(self) -> None:
        prepared = prepare(self.workbook, self.root / "prepared")
        by_name = {item["name"]: item for item in prepared["items"]}
        self.assertEqual(set(by_name), {"Claire's Korner", "Unlisted Place"})
        placeholder = by_name["Unlisted Place"]
        self.assertEqual(placeholder["metadata_status"], "review_only")
        self.assertIsNone(placeholder["average_cost"])
        self.assertIsNone(placeholder["open_after_8pm"])
        self.assertEqual(prepared["quality"]["aliases_applied"][0]["raw_rows"], 4)
        self.assertEqual(prepared["quality"]["aliases_applied"][0]["retained_rows"], 3)
        quarantined = [event for event in prepared["interactions"] if event["date"] is None]
        self.assertEqual(len(quarantined), 2)
        self.assertEqual({flag for event in quarantined for flag in event["quality_flags"]},
                         {"invalid_date", "missing_date"})
        self.assertTrue(any(event["item_id"] == placeholder["item_id"] for event in quarantined))

    def test_private_stable_hmac_ids_and_atomic_json_roundtrip(self) -> None:
        output = self.root / "prepared"
        first = prepare(self.workbook, output)
        second = prepare(self.workbook, output)
        self.assertEqual(first, second)
        self.assertEqual(first, load_prepared(output))
        secret_path = output / SECRET_FILENAME
        self.assertEqual(len(secret_path.read_bytes()), 32)
        self.assertEqual(stat.S_IMODE(secret_path.stat().st_mode), 0o600)
        different_secret = prepare(self.workbook, self.root / "other")
        self.assertNotEqual(first["interactions"][0]["user_id"],
                            different_secret["interactions"][0]["user_id"])
        self.assertEqual(first["items"], different_secret["items"])
        serialized = json.dumps(first)
        for token in ("PRIVATE_REVIEWER_SENTINEL", "PRIVATE_REVIEW_TEXT_SENTINEL",
                      "PRIVATE_DEMOGRAPHIC_SENTINEL", "DIFFERENT_PRIVATE_REVIEW_SENTINEL",
                      "SECOND_PRIVATE_REVIEWER", secret_path.read_bytes().hex()):
            self.assertFalse(token in serialized, "Sensitive source field or private secret leaked")

    def test_invalid_rating_fails_with_nonidentifying_error(self) -> None:
        make_workbook(self.workbook, invalid_rating=True)
        with self.assertRaises(DataValidationError) as caught:
            prepare(self.workbook, self.root / "prepared")
        self.assertEqual(str(caught.exception), "Invalid rating in review row 7")
        self.assertFalse((self.root / "prepared" / "prepared.json").exists())

    def test_invalid_existing_secret_is_rejected(self) -> None:
        output = self.root / "prepared"
        output.mkdir()
        (output / SECRET_FILENAME).write_bytes(b"invalid")
        with self.assertRaisesRegex(DataValidationError, "exactly 32 bytes"):
            prepare(self.workbook, output)


@unittest.skipUnless(source_workbook().is_file(), "Raw workbook not found at data/raw/")
class ActualWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary = tempfile.TemporaryDirectory()
        cls.output = Path(cls.temporary.name)
        cls.source = source_workbook()
        cls.before_hash = hashlib.sha256(cls.source.read_bytes()).hexdigest()
        cls.prepared = prepare(cls.source, cls.output)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temporary.cleanup()

    def test_actual_counts_and_anonymous_audits(self) -> None:
        quality = self.prepared["quality"]
        for field, expected in {
            "raw_catalog_rows": 63, "catalog_items": 63, "review_only_items": 3,
            "total_items": 66, "raw_review_rows": 1444,
            "exact_duplicate_rows_removed": 8, "retained_interactions": 1436,
            "valid_date_interactions": 1421, "quarantined_date_interactions": 15,
            "unique_users": 1047, "single_interaction_users": 960,
        }.items():
            self.assertEqual(quality[field], expected, field)
        self.assertEqual(sum(group["rows_removed"] for group in quality["exact_duplicate_audit"]), 8)
        self.assertEqual(quality["aliases_applied"], [{"from": "Clare's Korner",
                         "to": "Claire's Korner", "raw_rows": 1, "retained_rows": 1}])
        self.assertEqual(quality["catalog_value_repairs"], [{"restaurant": "Cross Rhodes",
                         "field": "longitude", "repair": "removed_stray_leading_comma"}])

    def test_unknown_items_never_disappear_and_references_are_complete(self) -> None:
        items = self.prepared["items"]
        self.assertEqual({item["name"] for item in items if item["metadata_status"] == "review_only"},
                         {"La Principal", "Todoroki Sushi", "World Market"})
        identifiers = {item["item_id"] for item in items}
        self.assertEqual(len(identifiers), 66)
        self.assertTrue(all(event["item_id"] in identifiers for event in self.prepared["interactions"]))
        self.assertEqual(len({event["event_id"] for event in self.prepared["interactions"]}), 1436)

    def test_invalid_dates_preserved_but_never_guessed(self) -> None:
        interactions = self.prepared["interactions"]
        self.assertEqual(sum(event["date"] is None for event in interactions), 15)
        for event in interactions:
            if event["date"] is None:
                self.assertEqual(event["quality_flags"], ["invalid_date"])
            else:
                datetime.strptime(event["date"], "%Y-%m-%d")
                self.assertEqual(event["quality_flags"], [])

    def test_source_is_unchanged_and_identity_text_demographics_are_absent(self) -> None:
        self.assertEqual(self.before_hash, hashlib.sha256(self.source.read_bytes()).hexdigest())
        serialized = json.dumps(self.prepared, ensure_ascii=False)
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        try:
            for row in list(workbook["Reviews"].iter_rows(values_only=True))[1:]:
                for column in (0, 3):
                    if isinstance(row[column], str) and row[column].strip():
                        token = json.dumps(row[column].strip(), ensure_ascii=False)
                        self.assertFalse(token in serialized, "Source identity or review text leaked")
        finally:
            workbook.close()
        allowed = {"event_id", "user_id", "item_id", "rating", "date", "source", "quality_flags"}
        self.assertTrue(all(set(event) == allowed for event in self.prepared["interactions"]))
        self.assertEqual(self.prepared, load_prepared(self.output / "prepared.json"))


if __name__ == "__main__":
    unittest.main()
