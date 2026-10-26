"""Ingest the raw ratings workbook into a validated, pseudonymized interaction log.

Only the Reviews and Restaurants worksheets are inputs; the other worksheets are
derived subsets, not additional feedback. Reviewer names are replaced by keyed
HMAC pseudonyms (RFC 2104), and review text and demographic columns never leave this
module. Invalid dates survive as quarantined events so that temporal evaluation
can exclude them explicitly.
"""

from __future__ import annotations

import hashlib
import hmac
import json
import math
import os
import re
import secrets
import tempfile
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REVIEW_HEADERS = (
    "Reviewer Name", "Restaurant Name", "Rating", "Review Text",
    "Date of Review", "Birth Year", "Marital Status", "Has Children?",
    "Vegetarian?", "Weight (lb)", "Height (in)", "Average Amount Spent",
    "Preferred Mode of Transport", "Northwestern Student?",
)
CATALOG_HEADERS = (
    "Restaurant Name", "Cuisine", "Latitude", "Longitude", "Average Cost",
    "Open After 8pm?", "Brief Description",
)
ALIASES = {"Clare's Korner": "Claire's Korner"}
SECRET_FILENAME = ".hmac_secret"
SCHEMA_VERSION = 2


class DataValidationError(ValueError):
    """A source row cannot be represented faithfully in the prepared schema."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _name(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value))


def _number(value: Any, field: str, row_number: int) -> float | None:
    if value is None or _text(value) == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise DataValidationError(f"Invalid {field} in catalog row {row_number}") from None
    if not math.isfinite(number):
        raise DataValidationError(f"Nonfinite {field} in catalog row {row_number}")
    return number


def _open_flag(value: Any, row_number: int) -> bool | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = _text(value).casefold()
    if normalized in {"yes", "true", "1"}:
        return True
    if normalized in {"no", "false", "0"}:
        return False
    raise DataValidationError(f"Invalid opening flag in catalog row {row_number}")


def _parse_date(value: Any) -> tuple[str | None, str | None]:
    """Accept actual calendar dates; never repair impossible dates by guessing."""
    if isinstance(value, datetime):
        return value.date().isoformat(), None
    if isinstance(value, date):
        return value.isoformat(), None
    if value is None or _text(value) == "":
        return None, "missing_date"
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%m-%d-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat(), None
            except ValueError:
                pass
    return None, "invalid_date"


def _canonical_cell(value: Any) -> list[Any]:
    """Preserve type and raw value when detecting exact *full-row* duplicates."""
    if isinstance(value, datetime):
        return ["datetime", value.isoformat()]
    if isinstance(value, date):
        return ["date", value.isoformat()]
    if value is None:
        return ["none", None]
    if isinstance(value, bool):
        return ["bool", value]
    if isinstance(value, (str, int, float)):
        return [type(value).__name__, value]
    return [type(value).__name__, str(value)]


def _row_bytes(row: tuple[Any, ...]) -> bytes:
    return json.dumps([_canonical_cell(v) for v in row], ensure_ascii=False,
                      separators=(",", ":"), allow_nan=False).encode("utf-8")


def _digest(secret: bytes, domain: str, value: bytes) -> str:
    return hmac.new(secret, domain.encode("ascii") + b"\0" + value,
                    hashlib.sha256).hexdigest()


def _private_secret(output_dir: Path) -> bytes:
    """Create once with mode 0600; preserve stable anonymous IDs on rebuilds."""
    path = output_dir / SECRET_FILENAME
    if path.is_symlink():
        raise DataValidationError("Private pseudonym secret must not be a symlink")
    try:
        fd = os.open(str(path), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError:
        secret = path.read_bytes()
    else:
        secret = secrets.token_bytes(32)
        with os.fdopen(fd, "wb") as handle:
            handle.write(secret)
    if len(secret) != 32:
        raise DataValidationError("Private pseudonym secret must contain exactly 32 bytes")
    os.chmod(path, 0o600)
    return secret


def _item_id(name: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.casefold()).strip("-") or "restaurant"
    return "restaurant-" + slug


def _write_json(path: Path, value: dict[str, Any]) -> None:
    """Atomic replacement avoids consumers observing a partially written file."""
    fd, temporary = tempfile.mkstemp(prefix="." + path.name + ".", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True,
                      allow_nan=False)
            handle.write("\n")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _rows(sheet: Any, expected: tuple[str, ...]) -> list[tuple[int, tuple[Any, ...]]]:
    values = sheet.iter_rows(values_only=True)
    try:
        header = next(values)
    except StopIteration:
        raise DataValidationError(f"Worksheet {sheet.title} is empty") from None
    if tuple(header[:len(expected)]) != expected:
        raise DataValidationError(f"Worksheet {sheet.title} has unexpected headers")
    return [(number, tuple(row)) for number, row in enumerate(values, 2)
            if any(value is not None for value in row)]


def prepare(workbook_path: str | Path, output_dir: str | Path) -> dict[str, Any]:
    """Validate, pseudonymize and save prepared.json plus anonymous quality data.

    This function does not write to the workbook. Exact full-row duplicates are
    the only review rows removed. Rows that cannot be faithfully represented
    fail with a field and row number, never their sensitive source values.
    """
    workbook_path, output_dir = Path(workbook_path), Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    secret = _private_secret(output_dir)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        missing = {"Reviews", "Restaurants"} - set(workbook.sheetnames)
        if missing:
            raise DataValidationError("Missing required Reviews or Restaurants worksheet")
        catalog_rows = _rows(workbook["Restaurants"], CATALOG_HEADERS)
        review_rows = _rows(workbook["Reviews"], REVIEW_HEADERS)
        ignored_sheets = [name for name in workbook.sheetnames
                          if name not in {"Reviews", "Restaurants"}]
    finally:
        workbook.close()

    items_by_name: dict[str, dict[str, Any]] = {}
    used_item_ids: set[str] = set()
    catalog_repairs: list[dict[str, str]] = []

    def add_item(name: str, item: dict[str, Any]) -> None:
        if name in items_by_name:
            raise DataValidationError("Duplicate normalized restaurant name in catalog")
        identifier = _item_id(name)
        if identifier in used_item_ids:
            identifier += "-" + hashlib.sha256(name.encode("utf-8")).hexdigest()[:10]
        used_item_ids.add(identifier)
        item.update(item_id=identifier, name=name)
        items_by_name[name] = item

    for row_number, row in catalog_rows:
        name = _name(row[0])
        if not name:
            raise DataValidationError(f"Missing restaurant name in catalog row {row_number}")
        latitude = _number(row[2], "latitude", row_number)
        raw_longitude = row[3]
        # One known spreadsheet typo has a stray comma before a numeric
        # longitude. This reversible syntax repair is explicit in the audit.
        comma_match = (re.fullmatch(r",\s*([+-]?\d+(?:\.\d+)?)", _text(raw_longitude))
                       if isinstance(raw_longitude, str) else None)
        if comma_match:
            raw_longitude = comma_match.group(1)
            catalog_repairs.append({"restaurant": name, "field": "longitude",
                                    "repair": "removed_stray_leading_comma"})
        longitude = _number(raw_longitude, "longitude", row_number)
        cost = _number(row[4], "average cost", row_number)
        if latitude is not None and not -90 <= latitude <= 90:
            raise DataValidationError(f"Out-of-range latitude in catalog row {row_number}")
        if longitude is not None and not -180 <= longitude <= 180:
            raise DataValidationError(f"Out-of-range longitude in catalog row {row_number}")
        if cost is not None and cost < 0:
            raise DataValidationError(f"Negative average cost in catalog row {row_number}")
        add_item(name, {
            "cuisine": _text(row[1]) or "Unknown",
            "latitude": latitude, "longitude": longitude, "average_cost": cost,
            "open_after_8pm": _open_flag(row[5], row_number),
            "description": _text(row[6]), "metadata_status": "catalog",
        })

    seen_rows: set[bytes] = set()
    duplicate_counts: Counter[str] = Counter()
    alias_raw_counts: Counter[str] = Counter()
    alias_retained_counts: Counter[str] = Counter()
    date_flags: Counter[str] = Counter()
    interactions: list[dict[str, Any]] = []
    raw_restaurant_names: set[str] = set()
    for row_number, row in review_rows:
        original_name = _name(row[1])
        raw_restaurant_names.add(original_name)
        if original_name in ALIASES:
            alias_raw_counts[original_name] += 1
        encoded_row = _row_bytes(row)
        event_id = "rating_" + _digest(secret, "event:v1", encoded_row)
        if encoded_row in seen_rows:
            duplicate_counts[event_id] += 1
            continue
        seen_rows.add(encoded_row)
        reviewer = _text(row[0])
        if not reviewer:
            raise DataValidationError(f"Missing reviewer identity in review row {row_number}")
        if not original_name:
            raise DataValidationError(f"Missing restaurant name in review row {row_number}")
        name = ALIASES.get(original_name, original_name)
        if original_name in ALIASES:
            alias_retained_counts[original_name] += 1
        if name not in items_by_name:
            add_item(name, {
                "cuisine": "Unknown", "latitude": None, "longitude": None,
                "average_cost": None, "open_after_8pm": None,
                # Left empty rather than filled with placeholder prose, which would leak into TF-IDF.
                "description": "",
                "metadata_status": "review_only",
            })
        raw_rating = row[2]
        try:
            rating = int(raw_rating)
        except (ValueError, TypeError, OverflowError):
            raise DataValidationError(f"Invalid rating in review row {row_number}") from None
        if isinstance(raw_rating, bool) or rating != raw_rating or not 1 <= rating <= 5:
            raise DataValidationError(f"Invalid rating in review row {row_number}")
        event_date, flag = _parse_date(row[4])
        if flag:
            date_flags[flag] += 1
        interactions.append({
            "event_id": event_id,
            "user_id": "u_" + _digest(secret, "user:v1", reviewer.encode("utf-8")),
            "item_id": items_by_name[name]["item_id"], "rating": rating,
            "date": event_date, "source": "explicit_rating",
            "quality_flags": [flag] if flag else [],
        })

    items = sorted(items_by_name.values(), key=lambda item: item["name"].casefold())
    user_counts = Counter(event["user_id"] for event in interactions)
    observed_pairs = {(event["user_id"], event["item_id"]) for event in interactions}
    dates = [event["date"] for event in interactions if event["date"] is not None]
    review_only = [item for item in items if item["metadata_status"] == "review_only"]
    source_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    provenance = {
        "workbook_filename": workbook_path.name, "workbook_sha256": source_sha256,
        "worksheets_used": ["Restaurants", "Reviews"],
        "worksheets_ignored": ignored_sheets,
        "identity": "Reviewer names are replaced by keyed HMAC-SHA256 pseudonyms. Names are not verified account IDs, so distinct people sharing a name merge.",
        "excluded_columns": "Reviewer name, review text and demographic columns are dropped at ingestion.",
        "catalog_snapshot": "The catalog is a single snapshot without per-item availability dates; offline candidate sets assume every item was available.",
    }
    quality = {
        "raw_catalog_rows": len(catalog_rows), "raw_review_rows": len(review_rows),
        "exact_duplicate_rows_removed": sum(duplicate_counts.values()),
        "retained_interactions": len(interactions), "catalog_items": len(catalog_rows),
        "review_only_items": len(review_only), "total_items": len(items),
        "raw_review_restaurant_names": len(raw_restaurant_names),
        "unique_users": len(user_counts),
        "unique_user_item_pairs": len(observed_pairs),
        "repeated_user_item_extra_rows": len(interactions) - len(observed_pairs),
        "single_interaction_users": sum(count == 1 for count in user_counts.values()),
        "interactions_per_user_histogram": {str(count): users for count, users in sorted(Counter(user_counts.values()).items())},
        "rating_histogram": {str(rating): count for rating, count in sorted(Counter(event["rating"] for event in interactions).items())},
        "valid_date_interactions": len(dates),
        "quarantined_date_interactions": sum(date_flags.values()),
        "date_quarantine_reason_counts": dict(sorted(date_flags.items())),
        "earliest_valid_date": min(dates) if dates else None,
        "latest_valid_date": max(dates) if dates else None,
        "matrix_observation_density": len(observed_pairs) / (len(user_counts) * len(items)) if user_counts and items else 0.0,
        "aliases_applied": [{"from": source, "to": target,
                             "raw_rows": alias_raw_counts[source],
                             "retained_rows": alias_retained_counts[source]}
                            for source, target in sorted(ALIASES.items())
                            if alias_raw_counts[source]],
        "exact_duplicate_audit": [{"event_id": event_id, "rows_removed": count}
                                  for event_id, count in sorted(duplicate_counts.items())],
        "review_only_restaurants": [{"item_id": item["item_id"], "name": item["name"],
                                     "missing_metadata": ["cuisine", "latitude", "longitude", "average_cost", "open_after_8pm", "description"]}
                                    for item in review_only],
        "catalog_value_repairs": catalog_repairs,
    }
    prepared = {"version": SCHEMA_VERSION, "provenance": provenance, "items": items,
                "interactions": interactions, "quality": quality}
    _write_json(output_dir / "data_quality.json", quality)
    _write_json(output_dir / "prepared.json", prepared)
    return prepared


def load_prepared(path: str | Path) -> dict[str, Any]:
    """Load a prepared JSON file or the output directory containing it."""
    path = Path(path)
    if path.is_dir():
        path = path / "prepared.json"
    with path.open(encoding="utf-8") as handle:
        prepared = json.load(handle)
    if prepared.get("version") != SCHEMA_VERSION:
        raise DataValidationError("Unsupported prepared data schema version")
    if not all(key in prepared for key in ("provenance", "items", "interactions", "quality")):
        raise DataValidationError("Prepared data is missing required sections")
    return prepared


def fingerprint(prepared: dict[str, Any]) -> str:
    """Content hash that ties trained artifacts to the exact prepared dataset."""
    return hashlib.sha256(json.dumps(prepared, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def public_quality(quality: dict[str, Any]) -> dict[str, Any]:
    """Aggregate data-quality stats without the per-event duplicate audit."""
    return {key: value for key, value in quality.items() if key != "exact_duplicate_audit"}
